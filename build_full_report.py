# -*- coding: utf-8 -*-
"""
Script para generar el Informe Final EP3 en formato DOCX
cumpliendo estrictamente con la pauta, rubrica y estilo del documento base.
"""
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    """Aplica color de fondo hexadecimal a una celda"""
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag.endswith('shd'):
            tcPr.remove(child)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Aplica margenes internos (padding) a una celda en twips"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def format_paragraph(p, font_name="Arial", size_pt=11, bold=False, italic=False, color_rgb=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_spacing=1.5, space_after=6, space_before=0):
    """Aplica formato formal a un parrafo"""
    p.alignment = align
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    for r in p.runs:
        r.font.name = font_name
        r.font.size = Pt(size_pt)
        r.bold = bold
        r.italic = italic
        if color_rgb:
            r.font.color.rgb = color_rgb

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(14)
    r.bold = True
    r.font.color.rgb = RGBColor(0, 51, 102)
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(12)
    r.bold = True
    r.font.color.rgb = RGBColor(0, 70, 130)
    return p

def add_heading_3(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(11)
    r.bold = True
    r.font.color.rgb = RGBColor(50, 50, 50)
    return p

def add_body_p(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.space_before = Pt(0)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Arial"
        r_pre.font.size = Pt(11)
        r_pre.bold = True
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(11)
    return p

def add_bullet_p(doc, bold_prefix, text):
    p = doc.add_paragraph(style='List Paragraph')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.space_before = Pt(0)
    r_pre = p.add_run(bold_prefix + ": ")
    r_pre.font.name = "Arial"
    r_pre.font.size = Pt(11)
    r_pre.bold = True
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(11)
    return p

print("Builder utilities defined successfully.")
